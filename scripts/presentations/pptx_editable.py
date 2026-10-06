"""Export the editable PowerPoint file from the deck and the frozen layout plan.

The exporter never receives HTML or a DOM.  Text, tables, shapes and connectors
are created through the native API so the recipient can edit them; an authored
block that has no native mapping fails with ``UNSUPPORTED_COMPONENT`` instead of
being flattened into a picture of the slide.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from scripts.checks.common import InputError
from scripts.presentations.ooxml import (
    CONNECTION_SITE, add_arrow_head, almost_transparent, box, emu, frame_margins, hide_from_slideshow,
    link_to_slide, link_to_url, set_description, set_notes,
)

SHAPES = {"rect": MSO_SHAPE.RECTANGLE, "round-rect": MSO_SHAPE.ROUNDED_RECTANGLE, "ellipse": MSO_SHAPE.OVAL}
TITLE_ONLY_LAYOUT = 5


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.lstrip("#").upper())


class EditableExporter:
    def __init__(self, document: dict[str, Any], layout: dict[str, Any]) -> None:
        self.document = document
        self.theme = document["theme"]
        self.layout = layout
        self.palette = self.theme["palette"]
        self.presentation = Presentation()
        self.presentation.slide_width = Emu(emu(layout["size_pt"]["width"]))
        self.presentation.slide_height = Emu(emu(layout["size_pt"]["height"]))
        self.slides: dict[str, Any] = {}
        self.objects: list[dict[str, Any]] = []
        self.exceptions: list[dict[str, str]] = []
        self.by_page: dict[str, list[dict[str, Any]]] = {}
        for element in layout["elements"]:
            self.by_page.setdefault(element["page_id"], []).append(element)

    # -- text --------------------------------------------------------------
    def typography(self, token: str) -> dict[str, Any]:
        return self.theme["typography"][token]

    def fill_frame(self, frame, lines: list[dict[str, Any]], token: str, *,
                   align: PP_ALIGN = PP_ALIGN.LEFT, indent: float = 0.0, marker: str | None = None) -> None:
        style = self.typography(token)
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 and not frame.paragraphs[0].runs else frame.add_paragraph()
            paragraph.alignment = align
            paragraph.line_spacing = Pt(style["leading"])
            paragraph.space_before = Pt(0)
            paragraph.space_after = Pt(0)
            if indent:
                properties = paragraph._p.get_or_add_pPr()
                properties.set("marL", str(emu(indent)))
                properties.set("indent", "0")
            pieces = list(line["runs"])
            if marker and index == 0:
                pieces = [{"text": f"{marker} ", "bold": False, "italic": False, "code": False}, *pieces]
            for piece in pieces:
                run = paragraph.add_run()
                run.text = piece["text"]
                run.font.size = Pt(style["size"])
                run.font.bold = bool(piece.get("bold") or style.get("bold"))
                run.font.italic = bool(piece.get("italic"))
                run.font.name = "IBM Plex Mono" if piece.get("code") or style["family"] == "mono" else "Carlito"
                run.font.color.rgb = rgb(self.palette[style["colour"]])

    def textbox(self, slide, element: dict[str, Any], token: str) -> Any:
        shape = slide.shapes.add_textbox(*box(element))
        frame_margins(shape.text_frame, self.theme["metrics"]["padding"] if token == "code" else 0.0)
        shape.text_frame.vertical_anchor = MSO_ANCHOR.TOP
        return shape

    # -- blocks ------------------------------------------------------------
    def compose(self, slide, page_id: str, element: dict[str, Any]) -> None:
        kind = element.get("type")
        if kind == "title":
            placeholder = slide.shapes.title
            left, top, width, height = box(element)
            placeholder.left, placeholder.top, placeholder.width, placeholder.height = left, top, width, height
            frame_margins(placeholder.text_frame)
            placeholder.text_frame.clear()
            self.fill_frame(placeholder.text_frame, element["lines"], "title")
            self.record(page_id, element, "title-placeholder")
            return
        if kind == "text":
            shape = self.textbox(slide, element, "text")
            shape.text_frame.clear()
            token = element.get("style_token") or "body"
            first = True
            for paragraph in element["paragraphs"]:
                marker = {"bullet": "\u2022", "number": "\u2013"}.get(paragraph["marker"] or "")
                self.fill_frame(shape.text_frame, paragraph["lines"], token,
                                indent=paragraph["indent"], marker=marker)
                first = False
            self.record(page_id, element, "text-frame")
            return
        if kind == "code":
            shape = self.textbox(slide, element, "code")
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(self.palette["panel"])
            shape.line.fill.background()
            shape.text_frame.clear()
            self.fill_frame(shape.text_frame, [{"runs": [{"text": line, "code": True}]} for line in element["lines"]], "code")
            self.record(page_id, element, "text-frame")
            return
        if kind == "table":
            rows = element["rows"]
            left, top, width, height = box(element)
            graphic = slide.shapes.add_table(len(rows), len(rows[0]["cells"]), left, top, width, height)
            table = graphic.table
            table.first_row = False
            table.horz_banding = False
            for index, column in enumerate(table.columns):
                column.width = Emu(emu(element["column_width"] + 2 * self.theme["metrics"]["padding"]))
            for row_index, row in enumerate(rows):
                table.rows[row_index].height = Emu(emu(row["height"]))
                for column_index, cell_lines in enumerate(row["cells"]):
                    cell = table.cell(row_index, column_index)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = rgb(self.palette[
                        "accent" if row["header"] else "panel" if row_index % 2 == 0 else "surface"])
                    frame_margins(cell.text_frame, self.theme["metrics"]["padding"])
                    cell.text_frame.clear()
                    self.fill_frame(cell.text_frame, cell_lines, "table-header" if row["header"] else "table")
            self.record(page_id, element, "table")
            return
        if kind == "image":
            picture = slide.shapes.add_picture(str(element["file"]), *box(element))
            set_description(picture, element["alt"])
            self.record(page_id, element, "picture")
            return
        if kind in ("shape", "diagram-node"):
            shape = slide.shapes.add_shape(SHAPES[element["shape"]], *box(element))
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(self.palette["panel"])
            shape.line.color.rgb = rgb(self.palette["line"])
            shape.text_frame.clear()
            frame_margins(shape.text_frame, self.theme["metrics"]["padding"])
            shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            self.fill_frame(shape.text_frame, element.get("lines", []), "shape", align=PP_ALIGN.CENTER)
            self.record(page_id, element, "autoshape")
            return
        if kind == "diagram-connector":
            self.connect(slide, page_id, element)
            return
        if kind in ("control", "chrome", "index-entry"):
            self.button(slide, page_id, element)
            return
        if kind == "diagram":
            return
        raise InputError(f"UNSUPPORTED_COMPONENT: {element['block_id']} has no native mapping for {kind!r}")

    def connect(self, slide, page_id: str, element: dict[str, Any]) -> None:
        points = element["points"]
        kind = MSO_CONNECTOR.STRAIGHT if element["kind"] == "straight" else MSO_CONNECTOR.ELBOW
        connector = slide.shapes.add_connector(
            kind, emu(points[0][0]), emu(points[0][1]), emu(points[-1][0]), emu(points[-1][1]))
        connector.line.color.rgb = rgb(self.palette["support"])
        connector.line.width = Pt(1.6)
        add_arrow_head(connector)
        start, end = self.shapes_by_block.get(element["from"]), self.shapes_by_block.get(element["to"])
        if start is None or end is None:
            raise InputError(f"connector {element['block_id']} has no anchored shapes")
        connector.begin_connect(start, CONNECTION_SITE[element["from_port"]])
        connector.end_connect(end, CONNECTION_SITE[element["to_port"]])
        self.record(page_id, element, "connector")

    def button(self, slide, page_id: str, element: dict[str, Any]) -> None:
        entry = element["type"] == "index-entry"
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, *box(element))
        shape.name = f"control::{element['block_id']}"
        set_description(shape, element["text"])
        enabled = element.get("enabled", True)
        colour = "accent" if element["type"] != "chrome" else "support"
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(self.palette["surface" if entry else
                                                     "panel" if not enabled else colour])
        shape.line.color.rgb = rgb(self.palette["line" if entry or not enabled else colour])
        shape.shadow.inherit = False
        shape.text_frame.clear()
        frame_margins(shape.text_frame, self.theme["metrics"]["padding"] / 2)
        shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        style = self.typography("control")
        paragraph = shape.text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.LEFT if entry else PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = element["text"]
        run.font.size = Pt(style["size"])
        run.font.bold = True
        run.font.name = "Carlito"
        run.font.color.rgb = rgb(self.palette["accent" if entry else
                                              "muted" if not enabled else "inverse"])
        self.pending.append((shape, page_id, element))
        self.record(page_id, element, "action-shape")

    def record(self, page_id: str, element: dict[str, Any], native: str) -> None:
        self.objects.append({"page_id": page_id, "block_id": element["block_id"], "native": native,
                             "text": element.get("text", "")})

    # -- assembly ----------------------------------------------------------
    def build(self, destination: Path) -> dict[str, Any]:
        navigation = {(item["page_id"], item["action_id"]): item for item in self.layout["navigation"]}
        layout = self.presentation.slide_layouts[TITLE_ONLY_LAYOUT]
        for page in self.layout["pages"]:
            slide = self.presentation.slides.add_slide(layout)
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = rgb(self.palette["surface"])
            self.slides[page["page_id"]] = slide
        self.pending: list[tuple[Any, str, dict[str, Any]]] = []
        for page in self.layout["pages"]:
            slide = self.slides[page["page_id"]]
            self.shapes_by_block: dict[str, Any] = {}
            elements = sorted(self.by_page.get(page["page_id"], []), key=lambda item: item["reading_order"])
            for element in elements:
                if element.get("type") == "diagram-connector":
                    continue
                before = len(slide.shapes._spTree)
                self.compose(slide, page["page_id"], element)
                if len(slide.shapes._spTree) > before:
                    self.shapes_by_block[element["block_id"]] = slide.shapes[-1]
            for element in elements:
                if element.get("type") == "diagram-connector":
                    self.compose(slide, page["page_id"], element)
            set_notes(slide, page["notes"])
            if page["kind"] == "support":
                hide_from_slideshow(slide)
        for shape, page_id, element in self.pending:
            record = navigation.get((page_id, element["action_id"]))
            if record is None:
                raise InputError(f"control {element['action_id']} on {page_id} has no planned destination")
            if record.get("reference_id"):
                link_to_url(shape, record["url"])
            elif record["enabled"]:
                link_to_slide(shape, self.slides[record["target_page_id"]])
        destination.parent.mkdir(parents=True, exist_ok=True)
        self.presentation.save(str(destination))
        return {"objects": self.objects, "exceptions": self.exceptions}


def export(document: dict[str, Any], layout: dict[str, Any], destination: Path) -> dict[str, Any]:
    return EditableExporter(document, layout).build(destination)
